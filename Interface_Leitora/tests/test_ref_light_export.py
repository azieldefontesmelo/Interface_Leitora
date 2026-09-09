from __future__ import annotations

import tempfile
import unittest
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from ref_light_export import append_ref_light_session


class RefLightExportTestCase(unittest.TestCase):
    def test_sessions_append_and_expand_reading_columns(self):
        with tempfile.TemporaryDirectory() as temporary_directory:
            path = Path(temporary_directory) / "ref_light.xlsx"
            first_date = datetime(2026, 9, 9, 10, 0, 0)
            second_date = datetime(2026, 9, 9, 10, 5, 0)

            append_ref_light_session(
                [10, 20],
                output_path=path,
                captured_at=first_date,
            )
            append_ref_light_session(
                [30, 40, 50],
                output_path=path,
                captured_at=second_date,
            )

            worksheet = load_workbook(path, data_only=True).active
            self.assertEqual(
                [worksheet.cell(1, column).value for column in range(1, 6)],
                ["Data", "Ref Light 1", "Ref Light 2", "Ref Light 3", "Média"],
            )
            self.assertEqual(worksheet.cell(2, 1).value, first_date)
            self.assertEqual(worksheet.cell(2, 2).value, 10)
            self.assertEqual(worksheet.cell(2, 3).value, 20)
            self.assertIsNone(worksheet.cell(2, 4).value)
            self.assertEqual(worksheet.cell(2, 5).value, 15)
            self.assertEqual(worksheet.cell(3, 1).value, second_date)
            self.assertEqual(
                [worksheet.cell(3, column).value for column in range(2, 6)],
                [30, 40, 50, 40],
            )


if __name__ == "__main__":
    unittest.main()
