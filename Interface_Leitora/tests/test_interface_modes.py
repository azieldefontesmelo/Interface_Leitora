from __future__ import annotations

import csv
import tempfile
import unittest
from datetime import datetime, timedelta
from pathlib import Path

from kivy.clock import Clock
from openpyxl import load_workbook

import interface_OSL
from database import Database, NEED_RE_READ_STATUS
from interface_OSL import AplicativoInterfaceOSL


class FakeSerial:
    def __init__(self):
        self.is_open = True
        self.writes = []
        self.in_waiting = 0
        self.flush_count = 0
        self.fail_writes = False

    def write(self, data):
        if self.fail_writes:
            raise interface_OSL.serial.SerialException("falha simulada")
        self.writes.append(data)
        return len(data)

    def flush(self):
        self.flush_count += 1

    def close(self):
        self.is_open = False


class InterfaceModeTestCase(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temporary_directory = tempfile.TemporaryDirectory()
        cls.root_path = Path(cls.temporary_directory.name)
        interface_OSL.TESTES_DIR = cls.root_path / "assets" / "testes"
        interface_OSL.LOG_SERIAL_DIR = cls.root_path / "assets" / "log"
        interface_OSL.SETTINGS_PATH = cls.root_path / "configuracoes.json"
        cls.database = Database(cls.root_path / "measurements.sqlite3")
        cls.database.register_dosimeter(
            "0123456789",
            ecc_hp10=1.25,
            ecc_hp007=1.5,
            bl_hp10=100,
            bl_hp007=200,
            begin_date="2025-01-01",
            end_date="2030-12-31",
        )
        cls.database.register_reader(
            "3001A01",
            rcf=0.000033,
            begin_date="2025-01-01",
        )
        cls.app = AplicativoInterfaceOSL()
        cls.app.database = cls.database
        cls.root = cls.app.build()
        cls.main = cls.root.get_screen("main")
        cls.bank = cls.root.get_screen("banco_dados")
        Clock.tick()
        cls.main.atualizar_leitoras_cadastradas()

    def setUp(self):
        self.database = Database(
            self.root_path / f"{self._testMethodName}.sqlite3"
        )
        self.database.register_dosimeter(
            "0123456789",
            ecc_hp10=1.25,
            ecc_hp007=1.5,
            bl_hp10=100,
            bl_hp007=200,
            begin_date="2025-01-01",
            end_date="2030-12-31",
        )
        self.database.register_reader(
            "3001A01",
            rcf=0.000033,
            begin_date="2025-01-01",
        )
        self.app.database = self.database
        self.main.database = self.database
        self.bank.database = self.database
        interface_OSL.TESTES_DIR = (
            self.root_path / self._testMethodName / "assets" / "testes"
        )
        interface_OSL.LOG_SERIAL_DIR = (
            self.root_path / self._testMethodName / "assets" / "log"
        )
        self.main.serial_connection = None
        self.main.log_arquivo = None
        self.main.current_measurement_id = None
        self.main.applied_parameters = None
        self.main._resetar_estado_alta_dose()
        self.main.last_dose_details = None
        self.main.fled_by_pled = dict(interface_OSL.FLED_PADRAO_POR_PLED)
        self.main.active_test_session_id = None
        self.main.active_test_dosimeter_id = None
        self.main.active_test_reading_type = None
        self.main.acquisition_active = False
        self.main.ref_light_mode_active = False
        self.main.ref_light_reading_active = False
        self.main.ref_light_readings = []
        self.main.ref_light_repetition_count = 0
        self.main.ref_light_target = 1
        self.main.ref_light_target_reached = False
        self.main.ref_light_average = 0
        self.main.test_session_active = False
        self.main.baseline_mode_active = False
        interface_OSL.SETTINGS_PATH.unlink(missing_ok=True)
        self.main.bl_update_mode = "MANUAL"
        self.main._sincronizar_botao_modo_bl()
        self.main.hp10_complete = False
        self.main.hp007_complete = False
        self.main.reading_type = "PERSONAL_DOSE"
        self.main.dose_channel = "HP10"
        self.main.test_mode = "MANUAL"
        self.main.ids.dosimeter_id_input.text = ""
        self.main._resetar_estado_bl()
        self.main._invalidar_dosimetro("Aguardando leitura do código de barras")
        self.main.atualizar_leitoras_cadastradas()

    def test_high_dose_flow_is_ordered_single_and_persisted(self):
        serial_port = FakeSerial()
        self.main.serial_connection = serial_port
        self.main.ids.nome_arquivo_input.text = "high-dose-flow.txt"
        self.main.ids.branco_textInput.text = "10"
        self.main.ids.rcf_textInput.text = "1"
        self.main.ids.ecc_textInput.text = "1"
        self.main.ids.fcal_textInput.text = "1"
        self.main.ids.fenerg_textInput.text = "1"

        self.main.botao_leitura()
        measurement_id = self.main.current_measurement_id
        self.main.processar_frame("#L1%AsatLeit")
        self.main.processar_frame("#L1%VsatLeit&")
        self.assertEqual(serial_port.writes, [b"#S1%SC1001&"])

        self.main.processar_frame(interface_OSL.FRAME_ALTA_DOSE)
        self.assertEqual(
            serial_port.writes[-1],
            interface_OSL.COMANDO_CONFIG_ALTA_DOSE.encode("ascii"),
        )
        self.assertEqual(
            self.main.high_dose_state,
            interface_OSL.ESTADO_ALTA_DOSE_AGUARDANDO_FILTRO,
        )
        self.assertIsNotNone(self.main.high_dose_popup)
        writes_after_first_trigger = list(serial_port.writes)
        self.main.processar_frame(interface_OSL.FRAME_ALTA_DOSE)
        self.assertEqual(serial_port.writes, writes_after_first_trigger)

        self.assertTrue(self.main._confirmar_filtro_alta_dose())
        self.assertFalse(self.main._confirmar_filtro_alta_dose())
        self.assertEqual(serial_port.writes[-1], b"#S1%SC1001&")
        self.assertEqual(serial_port.writes.count(b"#S1%SC1001&"), 2)
        self.main.processar_frame("#L1%A2&")
        self.assertEqual(
            self.main.high_dose_state,
            interface_OSL.ESTADO_LEITURA_ALTA_DOSE,
        )
        self.main.processar_frame("#L1%E45&")
        self.main.f_fechar_log = True
        self.main.processar_frame("#L1%D471&")

        measurement = self.database.get_measurement(measurement_id)
        self.assertEqual(measurement["dose_msv"], 190)
        self.assertIn("Alta dose; Fred=100", measurement["notes"])
        self.assertEqual(self.main.ids.label_dose.text, "190.000")
        self.assertEqual(self.main.last_dose_details["mode"], "Alta dose")
        self.assertEqual(self.main.high_dose_state, interface_OSL.ESTADO_LEITURA_NORMAL)
        self.assertGreaterEqual(serial_port.flush_count, 3)

    def test_high_dose_discards_samples_received_before_saturation(self):
        serial_port = FakeSerial()
        self.main.serial_connection = serial_port
        self.main.ids.nome_arquivo_input.text = "high-dose-discards-old-data.txt"
        self.main.ids.branco_textInput.text = "0"
        self.main.ids.rcf_textInput.text = "1"
        self.main.ids.ecc_textInput.text = "1"
        self.main.ids.fcal_textInput.text = "1"
        self.main.ids.fenerg_textInput.text = "1"

        self.main.botao_leitura()
        measurement_id = self.main.current_measurement_id
        self.main.processar_frame("#L1%A100&")
        self.main.processar_frame("#L1%E45&")
        self.main.processar_frame("#L1%D471&")
        self.assertEqual(self.main.soma, 100)
        caminho = Path(self.main.caminho_arquivo)
        self.assertIn("100", caminho.read_text(encoding="utf-8"))

        self.main.processar_frame(interface_OSL.FRAME_ALTA_DOSE)

        self.assertEqual(self.main.soma, 0)
        self.assertNotIn("100;45;471", caminho.read_text(encoding="utf-8"))
        self.assertEqual(self.main.high_dose_state, interface_OSL.ESTADO_ALTA_DOSE_AGUARDANDO_FILTRO)

        self.assertTrue(self.main._confirmar_filtro_alta_dose())
        self.main.processar_frame("#L1%A2&")
        self.main.processar_frame("#L1%E45&")
        self.main.f_fechar_log = True
        self.main.processar_frame("#L1%D471&")

        measurement = self.database.get_measurement(measurement_id)
        self.assertEqual(measurement["raw_signal"], 2)
        self.assertEqual(measurement["count_01s"], 2)
        self.assertNotIn("100;45;471", caminho.read_text(encoding="utf-8"))

    def test_high_dose_tx_failure_does_not_open_confirmation(self):
        serial_port = FakeSerial()
        self.main.serial_connection = serial_port
        self.main.ids.nome_arquivo_input.text = "high-dose-tx-error.txt"
        self.main.botao_leitura()
        measurement_id = self.main.current_measurement_id
        serial_port.fail_writes = True

        self.main.processar_frame(interface_OSL.FRAME_ALTA_DOSE)

        self.assertIsNone(self.main.high_dose_popup)
        self.assertEqual(self.main.high_dose_state, interface_OSL.ESTADO_LEITURA_NORMAL)
        self.assertEqual(self.database.get_measurement(measurement_id)["status"], "ERRO")

    def test_manual_high_dose_control_triggers_high_dose_flow(self):
        serial_port = FakeSerial()
        self.main.serial_connection = serial_port
        self.main.ids.nome_arquivo_input.text = "manual-high-dose.txt"
        self.main.botao_leitura()

        setup = self.root.get_screen("parametros")
        setup.ids.ler_alta_dose_button.dispatch("on_release")
        self.assertEqual(
            self.main.high_dose_state,
            interface_OSL.ESTADO_ALTA_DOSE_AGUARDANDO_FILTRO,
        )
        self.assertIsNotNone(self.main.high_dose_popup)
        self.assertEqual(
            serial_port.writes[-1],
            interface_OSL.COMANDO_CONFIG_ALTA_DOSE.encode("ascii"),
        )
        self.main.fechar_log(status="INTERROMPIDO", notes="teste")

    def test_transmitted_parameters_are_mirrored_in_setup(self):
        serial_port = FakeSerial()
        self.main.serial_connection = serial_port

        self.assertTrue(
            self.main.enviar_serial("#S1%M2G3L10000P1Z01000Q4&")
        )

        setup = self.root.get_screen("parametros")
        self.assertEqual(setup.ids.modo_input.text, "2")
        self.assertEqual(setup.ids.ganho_input.text, "3")
        self.assertEqual(setup.ids.tempo_leitura_input.text, "10000")
        self.assertEqual(setup.ids.potencia_input.text, "1")
        self.assertEqual(setup.ids.tempo_zeramento_input.text, "01000")
        self.assertEqual(setup.ids.potencia_zeramento_input.text, "4")

    def test_high_dose_completion_restores_stimulation_intensity_four(self):
        serial_port = FakeSerial()
        self.main.serial_connection = serial_port
        self.main.ids.nome_arquivo_input.text = "high-dose-restore.txt"

        self.main.botao_leitura()
        self.main.processar_frame(interface_OSL.FRAME_ALTA_DOSE)
        self.assertEqual(self.root.get_screen("parametros").ids.potencia_input.text, "1")
        self.assertTrue(self.main._confirmar_filtro_alta_dose())
        self.main.processar_frame("#L1%A2&")
        self.main.processar_frame("#L1%E45&")
        self.main.f_fechar_log = True
        self.main.processar_frame("#L1%D471&")

        setup = self.root.get_screen("parametros")
        self.assertEqual(setup.ids.potencia_input.text, "4")
        self.assertEqual(serial_port.writes[-1], b"#S1%M1G4L03000P4Z05000Q4&")

    def test_second_saturation_after_restart_finishes_as_error(self):
        serial_port = FakeSerial()
        self.main.serial_connection = serial_port
        self.main.ids.nome_arquivo_input.text = "persistent-saturation.txt"
        self.main.botao_leitura()
        measurement_id = self.main.current_measurement_id
        self.main.processar_frame(interface_OSL.FRAME_ALTA_DOSE)
        self.main._confirmar_filtro_alta_dose()
        self.main.processar_frame("#L1%A2&")
        self.main.processar_frame(interface_OSL.FRAME_ALTA_DOSE)

        measurement = self.database.get_measurement(measurement_id)
        self.assertEqual(measurement["status"], "ERRO")
        self.assertEqual(measurement["dose_msv"], 0)
        self.assertEqual(
            serial_port.writes.count(
                interface_OSL.COMANDO_CONFIG_ALTA_DOSE.encode("ascii")
            ),
            1,
        )

    def test_fled_is_loaded_edited_and_persisted_per_pled(self):
        setup = self.root.get_screen("parametros")
        setup.ids.potencia_input.text = "1"
        self.main.atualizar_fled_exibido("1")
        self.assertEqual(setup.ids.fled_input.text, "100")
        setup.ids.fled_input.text = "125.5"
        self.assertTrue(self.main.salvar_fled_atual())

        self.main.fled_by_pled = dict(interface_OSL.FLED_PADRAO_POR_PLED)
        self.main.carregar_configuracoes()
        self.assertEqual(self.main.fled_by_pled["1"], 125.5)
        self.assertIsNone(self.main.fled_by_pled["2"])
        self.assertEqual(self.main.fled_by_pled["4"], 1)

    def test_dose_format_has_one_rule(self):
        self.assertEqual(self.main.formatar_dose(0), "0")
        self.assertEqual(self.main.formatar_dose(0.067), "0.067")
        self.assertEqual(self.main.formatar_dose(12.3454), "12.345")

    def test_dose_details_popup_formats_dose_with_three_decimals(self):
        self.main.last_dose_details = {
            "mode": "Normal",
            "formula": "soma × RCF × ECC × Fang × Fenerg",
            "sum": 0,
            "fled": None,
            "baseline": 0,
            "rcf": 1,
            "ecc": 1,
            "fang": 1,
            "fenerg": 1,
            "dose": 0,
        }

        popup = self.main.mostrar_detalhes_dose()
        try:
            labels = [widget.text for widget in popup.content.children[1].children]
            self.assertIn("0.000 mSv", labels)
        finally:
            popup.dismiss()

    @classmethod
    def tearDownClass(cls):
        if cls.main.serial_aberta():
            cls.main.desconectar_serial(atualizar_botao=False)
        cls.temporary_directory.cleanup()

    def test_01_visual_mode_switch_and_focus(self):
        self.main.selecionar_modo("DOSIMETER_ID")
        Clock.tick()
        self.assertEqual(self.main.test_mode, "DOSIMETER_ID")
        self.assertEqual(self.main.ids.manual_panel.opacity, 0)
        self.assertEqual(self.main.ids.dosimeter_panel.opacity, 1)
        self.assertTrue(self.main.ids.dosimeter_id_input.focus)
        self.assertTrue(self.main.ids.start_button.disabled)

        self.main.selecionar_modo("MANUAL")
        Clock.tick()
        self.assertEqual(self.main.ids.manual_panel.opacity, 1)
        self.assertEqual(self.main.ids.dosimeter_panel.opacity, 0)
        self.assertFalse(self.main.ids.ecc_textInput.disabled)

    def test_hidden_mode_panel_does_not_block_text_fields(self):
        self.main.selecionar_modo("DOSIMETER_ID")
        Clock.tick()
        self.main.selecionar_modo("MANUAL")
        Clock.tick()

        manual_field = self.main.ids.rcf_textInput
        hidden_control = self.main.ids.reader_spinner
        self.assertTrue(hidden_control.collide_point(*manual_field.center))

        class Touch:
            pos = manual_field.center

        self.assertFalse(
            self.main.ids.dosimeter_panel.dispatch("on_touch_down", Touch())
        )

        manual_field.focus = True
        manual_field.text = ""
        manual_field.insert_text("0.000044")
        self.assertEqual(manual_field.text, "0.000044")

    def test_02_barcode_enter_validates_without_starting(self):
        self.main.selecionar_modo("DOSIMETER_ID")
        Clock.tick()
        self.main.ids.reader_spinner.text = "3001A01"
        self.main.ids.dosimeter_id_input.text = "\x020123456789\r\n"
        before = len(self.database.search_measurements())
        self.assertTrue(self.main.confirmar_codigo_dosimetro())
        Clock.tick()
        self.assertEqual(
            len(self.database.search_measurements()),
            before,
            "receber Enter não pode iniciar a aquisição",
        )
        self.assertTrue(self.main.start_allowed)
        self.assertFalse(self.main.ids.start_button.disabled)
        self.assertEqual(self.main.loaded_ecc, "1.25")
        self.assertEqual(self.main.loaded_bl, "100")
        self.assertEqual(self.main.loaded_rcf, "3.3e-05")
        self.assertTrue(
            self.main.automatic_file_name.startswith("0123456789_")
        )
        self.main.ids.arquivo_observacao_input.text = "sala 2"
        self.assertTrue(
            self.main.automatic_file_name.endswith("_sala 2.txt")
        )
        self.assertTrue(self.main.ids.dosimeter_id_input.focus)

        self.main.ids.dosimeter_id_input.text = "9999999999"
        self.assertFalse(self.main.confirmar_codigo_dosimetro())
        self.assertFalse(self.main.start_allowed)
        self.assertTrue(self.main.ids.start_button.disabled)

    def test_03_database_screen_crud_is_functional(self):
        self.bank.novo_dosimetro()
        self.bank.ids.db_dosimeter_id.text = "9876543210"
        self.bank.ids.db_dosimeter_ecc_hp10.text = "1,5"
        self.bank.ids.db_dosimeter_ecc_hp007.text = "1,7"
        self.bank.ids.db_dosimeter_bl_hp10.text = "150"
        self.bank.ids.db_dosimeter_bl_hp007.text = "170"
        self.bank.ids.db_dosimeter_begin.text = "01/01/2025"
        self.bank.ids.db_dosimeter_end.text = "31/12/2030"
        Clock.tick()
        self.bank.salvar_dosimetro()
        self.assertIn("sucesso", self.bank.dosimeter_message)
        self.assertEqual(
            self.database.get_dosimeter("9876543210")["ecc_hp10"],
            1.5,
        )
        self.assertEqual(
            self.database.get_dosimeter("9876543210")["ecc_hp007"], 1.7
        )
        self.assertEqual(
            self.database.get_dosimeter("9876543210")["bl_hp007"], 170
        )
        self.bank.alternar_dosimetro()
        self.assertFalse(
            self.database.get_dosimeter("9876543210")["active"]
        )

        self.bank.nova_leitora()
        self.bank.ids.db_reader_id.text = "READER-2"
        self.bank.ids.db_reader_rcf.text = "0,5"
        self.bank.ids.db_reader_begin.text = "01/01/2025"
        self.bank.ids.db_reader_end.text = ""
        Clock.tick()
        self.bank.salvar_leitora()
        self.assertIn("sucesso", self.bank.reader_message)
        self.assertEqual(self.database.get_reader("READER-2")["rcf"], 0.5)
        dosimeter_row = self.bank.ids.db_dosimeter_results.children[0]
        reader_row = self.bank.ids.db_reader_results.children[0]
        self.assertEqual(len(dosimeter_row.children), 8)
        self.assertEqual(len(reader_row.children), 5)
        self.assertIsNotNone(dosimeter_row.selection_callback)
        self.assertIsNotNone(reader_row.selection_callback)
        self.assertFalse(
            any("|" in cell.text for cell in dosimeter_row.children)
        )
        self.assertFalse(any("|" in cell.text for cell in reader_row.children))

        self.database.add_personal_dose(
            "0123456789",
            dose_dos=2.387,
        )
        self.database.add_background(
            "0123456789",
            counts=1912,
        )
        self.bank.pesquisar_doses_pessoais()
        self.bank.pesquisar_backgrounds()
        self.assertEqual(self.bank.personal_dose_count, "1 registros")
        self.assertEqual(self.bank.background_count, "1 registros")
        self.assertEqual(len(self.bank.ids.db_personal_dose_results.children), 1)
        self.assertEqual(len(self.bank.ids.db_background_results.children), 1)
        dataframe = self.bank._montar_dataframe_historico(
            self.database.search_personal_doses(),
            time_column="time_dos",
            hp10_column="hp10_dos",
            hp007_column="hp007_dos",
            status_column="status_dos",
        )
        self.assertEqual(
            list(dataframe.columns),
            [
                "Data/hora", "Dosímetro", "Hp(10) mSv",
                "Hp(0,07) mSv", "Status",
            ],
        )
        background_dataframe = self.bank._montar_dataframe_historico(
            self.database.search_backgrounds(),
            time_column="time_bg",
            hp10_column="hp10_counts",
            hp007_column="hp007_counts",
            status_column="status_bg",
            unit="Contagens",
        )
        self.assertEqual(
            list(background_dataframe.columns),
            [
                "Data/hora", "Dosímetro", "Hp(10) Contagens",
                "Hp(0,07) Contagens", "Status",
            ],
        )
        rendered_row = self.bank.ids.db_personal_dose_results.children[0]
        self.assertEqual(len(rendered_row.children), 5)
        self.assertFalse(
            any("|" in cell.text for cell in rendered_row.children)
        )

    def test_dosimeter_csv_export_includes_latest_bl_date(self):
        self.database.add_background(
            "0123456789",
            hp10_counts=1111,
            hp007_counts=1222,
            time_bg="2026-07-29T10:00:00Z",
        )
        self.database.add_background(
            "0123456789",
            hp10_counts=1333,
            hp007_counts=1444,
            time_bg="2026-07-30T11:00:00Z",
        )
        self.bank.ids.db_dosimeter_search.text = ""

        previous_assets_dir = interface_OSL.ASSETS_DIR
        interface_OSL.ASSETS_DIR = self.root_path / "csv_export"
        try:
            output = self.bank.exportar_csv_dosimetros()
        finally:
            interface_OSL.ASSETS_DIR = previous_assets_dir

        self.assertIsNotNone(output)
        with output.open(encoding="utf-8-sig", newline="") as file:
            rows = list(csv.DictReader(file))

        headers = list(
            self.bank._montar_dataframe_dosimetros(
                [],
                incluir_ultima_leitura_bl=True,
            ).columns
        )
        self.assertEqual(list(rows[0]), headers)
        self.assertEqual(rows[0][headers[0]], "0123456789")
        self.assertEqual(rows[0][headers[1]], "1.25")
        self.assertEqual(rows[0][headers[2]], "1.5")
        self.assertEqual(rows[0][headers[3]], "1333")
        self.assertEqual(rows[0][headers[4]], "1444")
        self.assertEqual(rows[0][headers[5]], "01/01/2025")
        self.assertEqual(rows[0][headers[6]], "31/12/2030")
        self.assertEqual(rows[0][headers[7]], "Ativo")
        expected_latest = datetime.fromisoformat(
            "2026-07-30T11:00:00+00:00"
        ).astimezone().strftime("%d/%m/%Y %H:%M:%S")
        self.assertEqual(rows[0][headers[8]], expected_latest)

    def test_history_csv_exports_match_screen_fields_with_date_last(self):
        self.database.add_personal_dose(
            "0123456789",
            hp10_dos=1.125,
            hp007_dos=2.25,
            time_dos="2026-07-30T11:00:00Z",
        )
        self.database.add_background(
            "0123456789",
            hp10_counts=1333,
            hp007_counts=1444,
            time_bg="2026-07-30T12:00:00Z",
        )

        previous_assets_dir = interface_OSL.ASSETS_DIR
        interface_OSL.ASSETS_DIR = self.root_path / "csv_history_export"
        try:
            personal_output = self.bank.exportar_csv_doses_pessoais()
            background_output = self.bank.exportar_csv_backgrounds()
        finally:
            interface_OSL.ASSETS_DIR = previous_assets_dir

        self.assertIsNotNone(personal_output)
        self.assertIsNotNone(background_output)
        with personal_output.open(encoding="utf-8-sig", newline="") as file:
            personal_rows = list(csv.DictReader(file))
        with background_output.open(encoding="utf-8-sig", newline="") as file:
            background_rows = list(csv.DictReader(file))

        personal_screen_columns = list(
            self.bank._montar_dataframe_historico(
                [],
                time_column="time_dos",
                hp10_column="hp10_dos",
                hp007_column="hp007_dos",
                status_column="status_dos",
            ).columns
        )
        background_screen_columns = list(
            self.bank._montar_dataframe_historico(
                [],
                time_column="time_bg",
                hp10_column="hp10_counts",
                hp007_column="hp007_counts",
                status_column="status_bg",
                unit="Contagens",
            ).columns
        )
        expected_personal_columns = [
            column for column in personal_screen_columns if column != "Data/hora"
        ] + ["Data/hora"]
        expected_background_columns = [
            column
            for column in background_screen_columns
            if column != "Data/hora"
        ] + ["Data/hora"]
        self.assertEqual(list(personal_rows[0]), expected_personal_columns)
        self.assertEqual(list(background_rows[0]), expected_background_columns)
        self.assertEqual(personal_rows[0][expected_personal_columns[0]], "0123456789")
        self.assertEqual(personal_rows[0][expected_personal_columns[1]], "1.125")
        self.assertEqual(personal_rows[0][expected_personal_columns[2]], "2.250")
        self.assertEqual(background_rows[0][expected_background_columns[0]], "0123456789")
        self.assertEqual(background_rows[0][expected_background_columns[1]], "1333")
        self.assertEqual(background_rows[0][expected_background_columns[2]], "1444")
        expected_personal_date = datetime.fromisoformat(
            "2026-07-30T11:00:00+00:00"
        ).astimezone().strftime("%d/%m/%Y %H:%M:%S")
        expected_background_date = datetime.fromisoformat(
            "2026-07-30T12:00:00+00:00"
        ).astimezone().strftime("%d/%m/%Y %H:%M:%S")
        self.assertEqual(
            personal_rows[0][expected_personal_columns[-1]], expected_personal_date
        )
        self.assertEqual(
            background_rows[0][expected_background_columns[-1]], expected_background_date
        )

    def test_csv_export_popup_can_export_only_todays_data(self):
        now = datetime.now()
        self.database.add_personal_dose(
            "0123456789",
            hp10_dos=1.125,
            hp007_dos=2.25,
            time_dos=now,
        )
        self.database.add_personal_dose(
            "0123456789",
            hp10_dos=3.125,
            hp007_dos=4.25,
            time_dos=now - timedelta(days=1),
        )

        previous_assets_dir = interface_OSL.ASSETS_DIR
        interface_OSL.ASSETS_DIR = self.root_path / "csv_today_export"
        popup = self.bank.abrir_popup_exportacao("personal")
        fields = self.bank._export_popup_fields
        try:
            fields["today"].active = True
            today_text = datetime.now().strftime("%d/%m/%Y")
            self.assertEqual(fields["initial"].text, today_text)
            self.assertEqual(fields["final"].text, today_text)
            self.assertTrue(fields["initial"].disabled)
            self.assertTrue(fields["final"].disabled)

            output = self.bank._confirmar_exportacao(
                popup,
                "personal",
                fields["initial"],
                fields["final"],
                fields["today"],
                fields["error"],
            )
        finally:
            if popup is not None:
                popup.dismiss()
            interface_OSL.ASSETS_DIR = previous_assets_dir

        self.assertIsNotNone(output)
        with output.open(encoding="utf-8-sig", newline="") as file:
            rows = list(csv.DictReader(file))
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["Hp(10) mSv"], "1.125")

    def test_dosimeter_csv_export_popup_filters_validity_period(self):
        self.database.register_dosimeter(
            "9876543210",
            ecc_hp10=1.1,
            ecc_hp007=1.2,
            begin_date="2031-01-01",
            end_date="2031-12-31",
        )

        previous_assets_dir = interface_OSL.ASSETS_DIR
        interface_OSL.ASSETS_DIR = self.root_path / "csv_dosimeter_export"
        popup = self.bank.abrir_popup_exportacao("dosimeters")
        fields = self.bank._export_popup_fields
        try:
            self.assertEqual(fields["initial"].text, "")
            self.assertEqual(fields["final"].text, "")
            fields["initial"].text = "01/01/2026"
            fields["final"].text = "31/12/2026"
            output = self.bank._confirmar_exportacao(
                popup,
                "dosimeters",
                fields["initial"],
                fields["final"],
                fields["today"],
                fields["error"],
            )
        finally:
            if popup is not None:
                popup.dismiss()
            interface_OSL.ASSETS_DIR = previous_assets_dir

        self.assertIsNotNone(output)
        with output.open(encoding="utf-8-sig", newline="") as file:
            rows = list(csv.DictReader(file))
        self.assertEqual([row["Dosímetro"] for row in rows], ["0123456789"])

    def test_unregistered_dosimeter_gets_defaults_and_tab_order(self):
        new_id = "0000000002"
        self.bank.ids.db_dosimeter_search.text = new_id
        self.bank.pesquisar_dosimetros()

        self.assertEqual(self.bank.ids.db_dosimeter_id.text, new_id)
        self.assertEqual(self.bank.ids.db_dosimeter_ecc_hp10.text, "1")
        self.assertEqual(self.bank.ids.db_dosimeter_ecc_hp007.text, "1")
        self.assertEqual(self.bank.ids.db_dosimeter_bl_hp10.text, "1")
        self.assertEqual(self.bank.ids.db_dosimeter_bl_hp007.text, "1")
        self.assertEqual(
            self.bank.ids.db_dosimeter_begin.text,
            datetime.now().strftime("%d/%m/%Y"),
        )
        self.assertEqual(self.bank.ids.db_dosimeter_end.text, "")

        focus_order = (
            "db_dosimeter_search",
            "db_dosimeter_id",
            "db_dosimeter_ecc_hp10",
            "db_dosimeter_ecc_hp007",
            "db_dosimeter_bl_hp10",
            "db_dosimeter_bl_hp007",
            "db_dosimeter_begin",
            "db_dosimeter_end",
        )
        for current_id, next_id in zip(focus_order, focus_order[1:] + focus_order[:1]):
            self.assertIs(
                self.bank.ids[current_id].focus_next,
                self.bank.ids[next_id],
            )

        self.bank.salvar_dosimetro()
        record = self.database.get_dosimeter(new_id)
        self.assertEqual(record["ecc_hp10"], 1)
        self.assertEqual(record["ecc_hp007"], 1)
        self.assertEqual(record["bl_hp10"], 1)
        self.assertEqual(record["bl_hp007"], 1)
        self.assertEqual(record["begin_date"], datetime.now().strftime("%Y-%m-%d"))

    def test_date_fields_do_not_insert_slashes_automatically(self):
        field = self.bank.ids.db_dosimeter_begin
        field.text = "21"
        Clock.tick()
        self.assertEqual(field.text, "21")
        field.text = "21/08/2026"
        Clock.tick()
        self.assertEqual(field.text, "21/08/2026")

    def test_exact_search_populates_edits_and_deletes_dosimeter(self):
        self.bank.ids.db_dosimeter_id.text = "0123456789"
        self.assertTrue(self.bank.carregar_dosimetro_por_id())
        self.assertEqual(self.bank.ids.db_dosimeter_id.text, "0123456789")
        self.assertEqual(self.bank.ids.db_dosimeter_ecc_hp10.text, "1.25")
        self.assertEqual(self.bank.ids.db_dosimeter_ecc_hp007.text, "1.5")
        self.assertEqual(self.bank.ids.db_dosimeter_bl_hp10.text, "100")
        self.assertEqual(self.bank.ids.db_dosimeter_bl_hp007.text, "200")
        self.assertFalse(self.bank.ids.db_dosimeter_id.disabled)

        self.bank.ids.db_dosimeter_id.text = "9876543210"
        self.bank.ids.db_dosimeter_ecc_hp10.text = "1,3"
        self.bank.ids.db_dosimeter_ecc_hp007.text = "1,6"
        self.bank.ids.db_dosimeter_bl_hp10.text = "110"
        self.bank.ids.db_dosimeter_bl_hp007.text = "210"
        self.bank.ids.db_dosimeter_begin.text = "02/02/2025"
        self.bank.ids.db_dosimeter_end.text = "02/02/2031"
        self.bank.salvar_dosimetro()
        self.assertIn("atualizado", self.bank.dosimeter_message)
        self.assertIsNone(self.database.get_dosimeter("0123456789"))
        edited = self.database.get_dosimeter("9876543210")
        self.assertEqual(edited["ecc_hp10"], 1.3)
        self.assertEqual(edited["bl_hp007"], 210)

        self.database.add_personal_dose(
            "9876543210",
            hp10_dos=1.1,
            hp007_dos=1.2,
        )
        self.assertTrue(
            self.bank._confirmar_exclusao_dosimetro(None, "9876543210")
        )
        self.assertIsNone(self.database.get_dosimeter("9876543210"))
        self.assertEqual(self.database.search_personal_doses(), [])

    def test_exact_search_populates_edits_and_deletes_reader(self):
        self.bank.ids.db_reader_id.text = "3001A01"
        self.assertTrue(self.bank.carregar_leitora_por_id())
        self.assertEqual(self.bank.ids.db_reader_id.text, "3001A01")
        self.assertEqual(self.bank.ids.db_reader_rcf.text, "3.3e-05")
        self.assertEqual(self.bank.ids.db_reader_begin.text, "01/01/2025")
        self.assertFalse(self.bank.ids.db_reader_id.disabled)

        self.bank.ids.db_reader_id.text = "READER-RENAMED"
        self.bank.ids.db_reader_rcf.text = "0,000044"
        self.bank.ids.db_reader_begin.text = "02/02/2025"
        self.bank.ids.db_reader_end.text = "02/02/2031"
        self.bank.salvar_leitora()
        self.assertIn("atualizada", self.bank.reader_message)
        self.assertIsNone(self.database.get_reader("3001A01"))
        edited = self.database.get_reader("READER-RENAMED")
        self.assertEqual(edited["rcf"], 0.000044)

        measurement_id = self.database.add_measurement(
            reader_id="READER-RENAMED",
            dosimeter_id="0123456789",
            test_mode="DOSIMETER_ID",
            reading_type="PERSONAL_DOSE",
            dose_channel="HP10",
            test_session_id="reader-delete-test",
            file_name="reader-delete.txt",
            status="CONCLUIDO",
        )
        self.assertTrue(
            self.bank._confirmar_exclusao_leitora(None, "READER-RENAMED")
        )
        self.assertIsNone(self.database.get_reader("READER-RENAMED"))
        self.assertIsNone(self.database.get_measurement(measurement_id))

    def test_filled_fields_validate_without_enter(self):
        self.main.selecionar_modo("DOSIMETER_ID")
        Clock.tick()
        self.main.ids.reader_spinner.text = "3001A01"
        self.main.ids.dosimeter_id_input.text = "0123456789"
        Clock.tick()

        self.assertTrue(self.main.start_allowed)
        self.assertFalse(self.main.ids.start_button.disabled)
        self.assertIn("pressione Start", self.main.dosimeter_status)

    def test_04_manual_acquisition_preserves_serial_and_history(self):
        self.main.selecionar_modo("MANUAL")
        self.main.serial_connection = FakeSerial()
        self.main.ids.ecc_textInput.text = "1"
        self.main.ids.rcf_textInput.text = "0.000033"
        self.main.ids.fcal_textInput.text = "1"
        self.main.ids.fenerg_textInput.text = "1"
        self.main.ids.branco_textInput.text = "0"
        self.main.ids.nome_arquivo_input.text = "manual-integration"

        self.main.botao_leitura()
        self.assertIsNotNone(self.main.current_measurement_id)
        measurement_id = self.main.current_measurement_id
        self.assertIn(interface_OSL.COMANDOS_SUDO["leitura"].encode("ascii"), self.main.serial_connection.writes)

        self.main.processar_frame("#L1%A1733")
        self.main.processar_frame("#L1%E45")
        self.main.f_fechar_log = True
        self.main.processar_frame("#L1%D471")

        record = self.database.get_measurement(measurement_id)
        self.assertEqual(record["test_mode"], "MANUAL")
        self.assertEqual(record["count_01s"], 1733)
        self.assertEqual(record["current_ma"], 45)
        self.assertEqual(record["light_mv"], 471)
        self.assertAlmostEqual(record["dose_msv"], 1733 * 0.000033)
        self.assertEqual(record["status"], "CONCLUIDO")
        self.assertTrue(Path(record["file_path"]).is_file())
        self.assertIsNone(self.main.current_measurement_id)
        self.bank.pesquisar_historico()
        history_row = self.bank.ids.db_history_results.children[0]
        self.assertEqual(len(history_row.children), 7)
        self.assertFalse(any("|" in cell.text for cell in history_row.children))
        history_row.selection_callback(history_row.record)
        self.assertIn(f"ID {measurement_id}", self.bank.history_details)

    def test_05_dosimeter_mode_uses_database_coefficients(self):
        self.main.selecionar_modo("DOSIMETER_ID")
        self.main.serial_connection = FakeSerial()
        self.main.atualizar_leitoras_cadastradas()
        self.main.ids.reader_spinner.text = "3001A01"
        self.main.ids.dosimeter_id_input.text = "0123456789"
        self.assertTrue(self.main.confirmar_codigo_dosimetro())

        self.main.botao_leitura()
        hp10_measurement_id = self.main.current_measurement_id
        self.assertIsNotNone(hp10_measurement_id)
        self.main.processar_frame("#L1%A1733")
        self.main.processar_frame("#L1%E45")
        self.main.f_fechar_log = True
        self.main.processar_frame("#L1%D471")

        hp10 = self.database.get_measurement(hp10_measurement_id)
        self.assertEqual(hp10["test_mode"], "DOSIMETER_ID")
        self.assertEqual(hp10["dosimeter_id"], "0123456789")
        self.assertEqual(hp10["reader_id"], "3001A01")
        self.assertEqual(hp10["reading_type"], "PERSONAL_DOSE")
        self.assertEqual(hp10["dose_channel"], "HP10")
        self.assertEqual(hp10["ecc_applied"], 1.25)
        self.assertEqual(hp10["baseline_applied"], 100)
        self.assertEqual(hp10["rcf_applied"], 0.000033)
        self.assertAlmostEqual(
            hp10["dose_msv"],
            (1733 - 100) * 0.000033 * 1.25,
        )
        self.assertEqual(self.database.search_personal_doses(), [])
        self.assertTrue(self.main.test_session_active)
        self.assertTrue(self.main.hp10_complete)
        self.assertEqual(self.main.dose_channel, "HP007")
        self.assertEqual(self.main.ids.dosimeter_id_input.text, "0123456789")

        self.main.botao_leitura()
        hp007_measurement_id = self.main.current_measurement_id
        self.assertEqual(
            self.main.ids.hp007_button.text,
            "Lendo Hp(0,07)...",
        )
        self.main.processar_frame("#L1%A2000")
        self.main.processar_frame("#L1%E45")
        self.main.f_fechar_log = True
        self.main.processar_frame("#L1%D471")

        hp007 = self.database.get_measurement(hp007_measurement_id)
        self.assertEqual(hp007["dose_channel"], "HP007")
        self.assertEqual(hp007["ecc_applied"], 1.5)
        self.assertEqual(hp007["baseline_applied"], 200)
        self.assertAlmostEqual(
            hp007["dose_msv"],
            (2000 - 200) * 0.000033 * 1.5,
        )
        history = self.database.search_personal_doses()
        self.assertEqual(len(history), 1)
        self.assertEqual(history[0]["hp10_measurement_id"], hp10_measurement_id)
        self.assertEqual(history[0]["hp007_measurement_id"], hp007_measurement_id)
        self.assertFalse(self.main.test_session_active)
        self.assertEqual(self.main.ids.dosimeter_id_input.text, "")
        Clock.tick()
        self.assertTrue(self.main.ids.dosimeter_id_input.focus)

    def test_052_channel_only_switches_after_completed_reading(self):
        self.main.selecionar_modo("DOSIMETER_ID")
        self.main.serial_connection = FakeSerial()
        self.main.ids.reader_spinner.text = "3001A01"
        self.main.ids.dosimeter_id_input.text = "0123456789"
        self.assertTrue(self.main.confirmar_codigo_dosimetro())

        self.main.botao_leitura()
        self.assertTrue(self.main.acquisition_active)
        self.assertEqual(self.main.ids.hp10_button.text, "Lendo Hp(10)...")
        self.assertIn("Lendo Hp(10)", self.main.dosimeter_status)
        self.assertTrue(self.main.ids.hp007_button.disabled)
        self.assertFalse(self.main.selecionar_grandeza("HP007"))
        self.assertEqual(self.main.dose_channel, "HP10")

        self.main.botao_stop()
        self.assertFalse(self.main.acquisition_active)
        self.assertTrue(self.main.ids.hp007_button.disabled)
        self.assertFalse(self.main.selecionar_grandeza("HP007"))
        self.assertEqual(self.main.dose_channel, "HP10")

        self.main.automatic_file_name = "retry-hp10.txt"
        self.main.botao_leitura()
        self.main.processar_frame("#L1%A1733")
        self.main.processar_frame("#L1%E45")
        self.main.f_fechar_log = True
        self.main.processar_frame("#L1%D471")

        self.assertEqual(self.main.dose_channel, "HP007")
        self.assertTrue(self.main.hp10_complete)
        self.assertEqual(self.main.ids.hp10_button.text, "Hp(10) concluído  ✓")
        self.assertFalse(self.main.ids.hp007_button.disabled)

    def test_high_dose_requires_new_reading_and_persists_tag(self):
        self.main.selecionar_modo("DOSIMETER_ID")
        self.main.serial_connection = FakeSerial()
        self.main.ids.reader_spinner.text = "3001A01"
        self.main.ids.dosimeter_id_input.text = "0123456789"
        self.assertTrue(self.main.confirmar_codigo_dosimetro())

        self.main.botao_leitura()
        self.main.processar_frame("#L1%A50000")
        self.main.processar_frame("#L1%E45")
        self.main.f_fechar_log = True
        self.main.processar_frame("#L1%D471")

        self.assertIsNone(self.main.re_read_popup)
        self.main.botao_leitura()
        self.main.processar_frame("#L1%A200")
        self.main.processar_frame("#L1%E45")
        self.main.f_fechar_log = True
        self.main.processar_frame("#L1%D471")

        history = self.database.search_personal_doses()
        self.assertEqual(len(history), 1)
        self.assertGreaterEqual(history[0]["dose_dos"], 2)
        self.assertEqual(history[0]["status_dos"], NEED_RE_READ_STATUS)
        popup = self.main.re_read_popup
        self.assertIsNotNone(popup)
        popup_text = "\n".join(
            child.text
            for child in popup.content.children
            if hasattr(child, "text")
        )
        self.assertIn("Refaça a leitura", popup_text)
        self.assertIn(NEED_RE_READ_STATUS, popup_text)
        popup.dismiss()

    def test_low_dose_asks_to_save_baseline_counts(self):
        self.main.selecionar_modo("DOSIMETER_ID")
        self.main.serial_connection = FakeSerial()
        self.main.ids.reader_spinner.text = "3001A01"
        self.main.ids.dosimeter_id_input.text = "0123456789"
        self.assertTrue(self.main.confirmar_codigo_dosimetro())

        self.main.botao_leitura()
        self.main.processar_frame("#L1%A101")
        self.main.processar_frame("#L1%E45")
        self.main.f_fechar_log = True
        self.main.processar_frame("#L1%D471")

        self.main.botao_leitura()
        self.main.processar_frame("#L1%A201")
        self.main.processar_frame("#L1%E45")
        self.main.f_fechar_log = True
        self.main.processar_frame("#L1%D471")

        history = self.database.search_personal_doses()
        self.assertEqual(len(history), 1)
        self.assertEqual(history[0]["status_dos"], "Need to Erase")
        popup = self.main.baseline_save_popup
        self.assertIsNotNone(popup)
        popup_text = "\n".join(
            child.text
            for child in popup.content.children
            if hasattr(child, "text")
        )
        self.assertIn("salvar esta leitura como baseline", popup_text)

        self.assertTrue(
            self.main._salvar_dose_como_baseline(popup, history[0])
        )
        background = self.database.get_latest_background("0123456789")
        self.assertEqual(background["hp10_counts"], 101)
        self.assertEqual(background["hp007_counts"], 201)
        self.assertEqual(self.database.get_dosimeter("0123456789")["bl_hp10"], 101)
        self.assertEqual(self.database.get_dosimeter("0123456789")["bl_hp007"], 201)
        self.assertIsNone(self.main.baseline_save_popup)

    def test_low_hp007_dose_also_asks_to_save_baseline_counts(self):
        self.main.selecionar_modo("DOSIMETER_ID")
        self.main.serial_connection = FakeSerial()
        self.main.ids.reader_spinner.text = "3001A01"
        self.main.ids.dosimeter_id_input.text = "0123456789"
        self.assertTrue(self.main.confirmar_codigo_dosimetro())

        self.main.botao_leitura()
        self.main.processar_frame("#L1%A1000")
        self.main.processar_frame("#L1%E45")
        self.main.f_fechar_log = True
        self.main.processar_frame("#L1%D471")

        self.main.botao_leitura()
        self.main.processar_frame("#L1%A201")
        self.main.processar_frame("#L1%E45")
        self.main.f_fechar_log = True
        self.main.processar_frame("#L1%D471")

        history = self.database.search_personal_doses()[0]
        self.assertGreater(history["hp10_dos"], 0.01)
        self.assertLess(history["hp007_dos"], 0.01)
        self.assertEqual(history["status_dos"], "Need to Erase")
        self.assertIsNotNone(self.main.baseline_save_popup)
        popup_text = "\n".join(
            child.text
            for child in self.main.baseline_save_popup.content.children
            if hasattr(child, "text")
        )
        self.assertIn("Hp(0,07)", popup_text)
        self.main.baseline_save_popup.dismiss()

    def test_high_hp007_dose_also_requires_new_reading(self):
        self.main.selecionar_modo("DOSIMETER_ID")
        self.main.serial_connection = FakeSerial()
        self.main.ids.reader_spinner.text = "3001A01"
        self.main.ids.dosimeter_id_input.text = "0123456789"
        self.assertTrue(self.main.confirmar_codigo_dosimetro())

        self.main.botao_leitura()
        self.main.processar_frame("#L1%A1000")
        self.main.processar_frame("#L1%E45")
        self.main.f_fechar_log = True
        self.main.processar_frame("#L1%D471")

        self.main.botao_leitura()
        self.main.processar_frame("#L1%A50000")
        self.main.processar_frame("#L1%E45")
        self.main.f_fechar_log = True
        self.main.processar_frame("#L1%D471")

        history = self.database.search_personal_doses()[0]
        self.assertLess(history["hp10_dos"], 2)
        self.assertGreaterEqual(history["hp007_dos"], 2)
        self.assertEqual(history["status_dos"], NEED_RE_READ_STATUS)
        self.assertIsNotNone(self.main.re_read_popup)
        popup_text = "\n".join(
            child.text
            for child in self.main.re_read_popup.content.children
            if hasattr(child, "text")
        )
        self.assertIn("Hp(0,07)", popup_text)
        self.main.re_read_popup.dismiss()

    def test_ref_light_repetitions_are_averaged_and_exported_to_xlsx(self):
        xlsx_path = self.root_path / self._testMethodName / "documentos" / "ref_light.xlsx"
        previous_xlsx_path = interface_OSL.REF_LIGHT_XLSX_PATH
        interface_OSL.REF_LIGHT_XLSX_PATH = xlsx_path
        try:
            self.main.serial_connection = FakeSerial()
            self.assertTrue(self.main.botao_ref_light())
            self.main.ids.ref_light_repetition_input.text = "120"
            self.assertEqual(self.main.ref_light_target, 120)
            self.main.ids.ref_light_repetition_input.text = "3"
            self.assertEqual(self.main.ref_light_target, 3)

            for value in (100, 200, 300):
                self.main.botao_leitura()
                self.assertTrue(self.main.ref_light_reading_active)
                self.main.f_fechar_log = True
                self.main.processar_frame(f"#L1%D{value}")

            self.assertFalse(self.main.ref_light_mode_active)
            self.assertEqual(self.main.ref_light_readings, [100.0, 200.0, 300.0])
            self.assertEqual(self.main.ref_light_average, 200.0)
            self.assertTrue(xlsx_path.is_file())

            worksheet = load_workbook(xlsx_path, data_only=True).active
            self.assertEqual(
                [worksheet.cell(1, column).value for column in range(1, 6)],
                ["Data", "Ref Light 1", "Ref Light 2", "Ref Light 3", "Média"],
            )
            self.assertEqual(
                [worksheet.cell(2, column).value for column in range(2, 6)],
                [100.0, 200.0, 300.0, 200.0],
            )
            self.assertIsInstance(worksheet.cell(2, 1).value, datetime)
        finally:
            interface_OSL.REF_LIGHT_XLSX_PATH = previous_xlsx_path

    def test_ref_light_does_not_require_dosimeter_and_finishes_on_reader_frame(self):
        xlsx_path = self.root_path / self._testMethodName / "ref_light.xlsx"
        previous_xlsx_path = interface_OSL.REF_LIGHT_XLSX_PATH
        interface_OSL.REF_LIGHT_XLSX_PATH = xlsx_path
        try:
            self.main.selecionar_modo("DOSIMETER_ID")
            self.main.serial_connection = FakeSerial()
            self.assertFalse(self.main.start_allowed)
            self.assertTrue(self.main.botao_ref_light())
            self.assertFalse(self.main.ids.start_button.disabled)

            self.main.botao_leitura()
            self.assertTrue(self.main.ref_light_reading_active)
            self.main.processar_frame("#L1%D321")
            self.main.processar_frame("#L1%I0000000")

            self.assertFalse(self.main.ref_light_reading_active)
            self.assertFalse(self.main.ref_light_mode_active)
            self.assertEqual(self.main.ref_light_readings, [321.0])
            self.assertTrue(xlsx_path.is_file())
        finally:
            interface_OSL.REF_LIGHT_XLSX_PATH = previous_xlsx_path

    def test_ref_light_export_error_is_not_reported_as_finalizing(self):
        self.main.serial_connection = FakeSerial()
        self.assertTrue(self.main.botao_ref_light())
        self.main.ids.ref_light_repetition_input.text = "1"
        self.main.botao_leitura()
        self.main.f_fechar_log = True

        previous_exporter = interface_OSL.append_ref_light_session
        interface_OSL.append_ref_light_session = lambda *args, **kwargs: (
            (_ for _ in ()).throw(ImportError("openpyxl ausente"))
        )
        try:
            self.main.processar_frame("#L1%D321")
            self.assertTrue(self.main.ref_light_mode_active)
            self.assertIn("openpyxl ausente", self.main.dosimeter_status)
            self.assertNotIn("finalizando", self.main.dosimeter_status)
        finally:
            interface_OSL.append_ref_light_session = previous_exporter

    def test_055_post_erase_reading_is_saved_as_background(self):
        self.main.selecionar_modo("DOSIMETER_ID")
        self.main.serial_connection = FakeSerial()
        self.assertEqual(self.main.bl_update_mode, "MANUAL")
        self.assertEqual(
            self.main.ids.grandeza_side.width,
            self.main.ids.bl_repetition_controls.width,
        )
        self.assertEqual(self.main.ids.bl_repetition_controls.opacity, 0)
        self.assertTrue(self.main.ids.bl_repetition_controls.disabled)
        self.main.botao_apagar()
        self.assertTrue(self.main.baseline_mode_active)
        self.assertEqual(self.main.ids.erase_button.text, "Finalizar modo BL")
        self.assertEqual(self.main.ids.bl_repetition_controls.opacity, 1)
        self.assertFalse(self.main.ids.bl_repetition_controls.disabled)
        self.main.ids.bl_repetition_input.text = "3"
        self.main.ids.reader_spinner.text = "3001A01"
        self.main.ids.dosimeter_id_input.text = "0123456789"
        self.assertTrue(self.main.confirmar_codigo_dosimetro())

        hp10_ids = []
        for counts in (1733, 1744, 1755):
            self.main.botao_leitura()
            hp10_ids.append(self.main.current_measurement_id)
            self.main.processar_frame(f"#L1%A{counts}")
            self.main.processar_frame("#L1%E45")
            self.main.f_fechar_log = True
            self.main.processar_frame("#L1%D471")
            self.assertEqual(self.main.dose_channel, "HP10")

        self.assertEqual(self.main.bl_hp10_count, 3)
        self.assertTrue(self.main.bl_target_reached)
        self.assertTrue(self.main.selecionar_grandeza("HP007"))
        self.assertEqual(self.main.dose_channel, "HP007")
        self.main.ids.bl_repetition_input.text = "2"

        hp007_ids = []
        for counts in (1811, 1822):
            self.main.botao_leitura()
            hp007_ids.append(self.main.current_measurement_id)
            self.main.processar_frame(f"#L1%A{counts}")
            self.main.processar_frame("#L1%E45")
            self.main.f_fechar_log = True
            self.main.processar_frame("#L1%D471")
            self.assertEqual(self.main.dose_channel, "HP007")

        session_id = self.main.active_test_session_id
        readings = self.database.get_baseline_session_measurements(session_id)
        self.assertEqual(len(readings), 5)
        for record in readings:
            self.assertEqual(record["reading_type"], "BACKGROUND")
            self.assertEqual(record["dose_msv"], 0)
            self.assertEqual(record["ecc_applied"], 1)
            self.assertEqual(record["rcf_applied"], 1)
            self.assertEqual(record["baseline_applied"], 0)
        before_selection = self.database.get_dosimeter("0123456789")
        self.assertEqual(before_selection["bl_hp10"], 100)
        self.assertEqual(before_selection["bl_hp007"], 200)
        self.assertIsNone(self.database.get_latest_background("0123456789"))

        self.assertTrue(self.main.abrir_selecao_bl())
        self.assertIsNotNone(self.main.bl_selection_popup)
        self.main._selecionar_leitura_bl("HP10", hp10_ids[1], "down")
        self.main._selecionar_leitura_bl("HP007", hp007_ids[0], "down")
        background = self.main.aplicar_selecao_bl()
        self.assertEqual(background["hp10_counts"], 1744)
        self.assertEqual(background["hp007_counts"], 1811)
        dosimeter = self.database.get_dosimeter("0123456789")
        self.assertEqual(dosimeter["bl_hp10"], 1744)
        self.assertEqual(dosimeter["bl_hp007"], 1811)
        self.assertFalse(self.main.baseline_mode_active)
        self.assertEqual(self.main.reading_type, "PERSONAL_DOSE")
        self.assertEqual(self.main.ids.bl_repetition_controls.opacity, 0)
        self.assertTrue(self.main.ids.bl_repetition_controls.disabled)

    def test_056_automatic_bl_mode_applies_latest_channel_readings(self):
        self.assertEqual(self.main.alternar_modo_bl(), "AUTOMATICO")
        mode_button = self.root.get_screen("parametros").ids.bl_mode_button
        self.assertIn("Automático", mode_button.text)
        self.main.bl_update_mode = "MANUAL"
        self.assertEqual(self.main.carregar_configuracoes(), "AUTOMATICO")

        self.main.selecionar_modo("DOSIMETER_ID")
        self.main.serial_connection = FakeSerial()
        self.main.botao_apagar()
        self.main.ids.bl_repetition_input.text = "2"
        self.main.ids.reader_spinner.text = "3001A01"
        self.main.ids.dosimeter_id_input.text = "0123456789"
        self.assertTrue(self.main.confirmar_codigo_dosimetro())

        for counts in (1733, 1744):
            self.main.botao_leitura()
            self.main.processar_frame(f"#L1%A{counts}")
            self.main.processar_frame("#L1%E45")
            self.main.f_fechar_log = True
            self.main.processar_frame("#L1%D471")

        self.assertTrue(self.main.selecionar_grandeza("HP007"))
        self.main.ids.bl_repetition_input.text = "2"
        for counts in (1811, 1822):
            self.main.botao_leitura()
            self.main.processar_frame(f"#L1%A{counts}")
            self.main.processar_frame("#L1%E45")
            self.main.f_fechar_log = True
            self.main.processar_frame("#L1%D471")

        self.main.botao_apagar()
        dosimeter = self.database.get_dosimeter("0123456789")
        self.assertEqual(dosimeter["bl_hp10"], 1744)
        self.assertEqual(dosimeter["bl_hp007"], 1822)
        background = self.database.get_latest_background("0123456789")
        self.assertEqual(background["hp10_counts"], 1744)
        self.assertEqual(background["hp007_counts"], 1822)
        self.assertFalse(self.main.baseline_mode_active)
        self.assertIsNone(self.main.bl_selection_popup)

    def test_automatic_bl_displays_dose_but_persists_counts(self):
        self.main.ids.fcal_textInput.text = "1"
        self.main.ids.fenerg_textInput.text = "1"
        self.main.alternar_modo_bl()
        self.main.selecionar_modo("DOSIMETER_ID")
        self.main.serial_connection = FakeSerial()
        self.main.botao_apagar()
        self.main.ids.reader_spinner.text = "3001A01"
        self.main.ids.dosimeter_id_input.text = "0123456789"
        self.assertTrue(self.main.confirmar_codigo_dosimetro())

        self.main.botao_leitura()
        measurement_id = self.main.current_measurement_id
        self.assertEqual(self.main.ids.LabelDose.text, "Dose (mSv)")
        self.main.processar_frame("#L1%A1733")
        self.main.processar_frame("#L1%E45")
        self.main.f_fechar_log = True
        self.main.processar_frame("#L1%D471")

        expected_dose = self.main.formatar_dose(1733 * 0.000033 * 1.25)
        self.assertEqual(self.main.ids.label_dose.text, expected_dose)
        record = self.database.get_measurement(measurement_id)
        self.assertEqual(record["reading_type"], "BACKGROUND")
        self.assertEqual(record["raw_signal"], 1733)
        self.assertEqual(record["dose_msv"], 0)

    def test_06_stop_marks_measurement_as_interrupted(self):
        self.main.selecionar_modo("MANUAL")
        self.main.serial_connection = FakeSerial()
        self.main.ids.nome_arquivo_input.text = "manual-interrupted"
        self.main.ids.branco_textInput.text = "0"
        self.main.botao_leitura()
        measurement_id = self.main.current_measurement_id
        self.main.botao_stop()

        record = self.database.get_measurement(measurement_id)
        self.assertEqual(record["status"], "INTERROMPIDO")
        self.assertIn(
            interface_OSL.COMANDOS_SUDO["stop"].encode("ascii"),
            self.main.serial_connection.writes,
        )

    def test_missing_first_frame_does_not_leave_acquisition_stuck(self):
        self.main.selecionar_modo("MANUAL")
        self.main.serial_connection = FakeSerial()
        self.main.ids.nome_arquivo_input.text = "serial-timeout"
        self.main.ids.branco_textInput.text = "0"
        self.main.botao_leitura()
        measurement_id = self.main.current_measurement_id

        self.main._timeout_primeiro_frame(0)

        record = self.database.get_measurement(measurement_id)
        self.assertEqual(record["status"], "ERRO")
        self.assertIn("Nenhum frame recebido", record["notes"])
        self.assertFalse(self.main.acquisition_active)
        self.assertIsNone(self.main.log_arquivo)
        self.assertIn(
            interface_OSL.COMANDOS_SUDO["stop"].encode("ascii"),
            self.main.serial_connection.writes,
        )

    def test_equipment_end_frame_finishes_reading(self):
        self.main.selecionar_modo("MANUAL")
        self.main.serial_connection = FakeSerial()
        self.main.ids.nome_arquivo_input.text = "serial-end-frame"
        self.main.ids.branco_textInput.text = "0"
        self.main.botao_leitura()
        measurement_id = self.main.current_measurement_id
        self.main.f_fechar_log = False

        self.main.processar_frame("#L1%A1733")
        self.main.processar_frame("#L1%E45")
        self.main.processar_frame("#L1%D471")
        self.main.processar_frame("#L1%I0010000")

        record = self.database.get_measurement(measurement_id)
        self.assertEqual(record["status"], "CONCLUIDO")
        self.assertIn("Fim sinalizado", record["notes"])
        self.assertIsNone(self.main.current_measurement_id)


if __name__ == "__main__":
    unittest.main()
