from __future__ import annotations

from datetime import datetime
from pathlib import Path
from statistics import fmean

from app_paths import USER_DATA_DIR


REF_LIGHT_XLSX_PATH = USER_DATA_DIR / "ref_light.xlsx"


def append_ref_light_session(
    readings,
    *,
    output_path: str | Path | None = None,
    captured_at: datetime | None = None,
) -> tuple[Path, float]:
    """Append one Ref Light session to the persistent XLSX workbook.

    Each session occupies one row. The first column is the capture date, the
    following columns contain the individual readings, and the final column
    contains the session average. The workbook grows its reading columns if
    a later session has more repetitions than previous sessions.
    """

    values = [float(value) for value in readings]
    if not values:
        raise ValueError("A sessão Ref Light precisa de pelo menos uma leitura")

    # openpyxl is intentionally imported lazily so the rest of the reader can
    # still start when an installation is missing the optional XLSX engine.
    try:
        from openpyxl import Workbook, load_workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError as error:
        raise RuntimeError(
            "A exportação Ref Light exige a dependência openpyxl. "
            "Instale as dependências do requirements.txt."
        ) from error

    path = Path(output_path or REF_LIGHT_XLSX_PATH)
    path.parent.mkdir(parents=True, exist_ok=True)
    average = float(fmean(values))
    timestamp = captured_at or datetime.now()
    required_reading_columns = len(values)

    if path.is_file():
        workbook = load_workbook(path)
        worksheet = workbook.active
    else:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Ref Light"

    existing_reading_columns = _reading_column_count(worksheet)
    reading_columns = max(existing_reading_columns, required_reading_columns)
    _ensure_headers(
        worksheet,
        reading_columns,
        required_reading_columns,
        Font,
        PatternFill,
        Alignment,
    )

    row_number = worksheet.max_row + 1
    worksheet.cell(row=row_number, column=1, value=timestamp)
    worksheet.cell(row=row_number, column=1).number_format = "dd/mm/yyyy hh:mm:ss"
    for index, value in enumerate(values, start=2):
        worksheet.cell(row=row_number, column=index, value=value)
    worksheet.cell(row=row_number, column=reading_columns + 2, value=average)

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = (
        f"A1:{_column_letter(reading_columns + 2)}{worksheet.max_row}"
    )
    worksheet.column_dimensions["A"].width = 21
    for index in range(2, reading_columns + 2):
        worksheet.column_dimensions[_column_letter(index)].width = 15

    # Replace the file atomically so a failed write does not destroy previous
    # sessions. Excel may refuse the replacement while the workbook is open;
    # that error is deliberately returned to the UI for the operator.
    temporary_path = path.with_suffix(".tmp.xlsx")
    try:
        workbook.save(temporary_path)
        temporary_path.replace(path)
    finally:
        if temporary_path.exists():
            temporary_path.unlink()

    return path, average


def _reading_column_count(worksheet) -> int:
    if worksheet.max_row < 1 or worksheet.max_column < 2:
        return 0
    headers = [worksheet.cell(row=1, column=index).value for index in range(2, worksheet.max_column + 1)]
    return sum(1 for header in headers if str(header or "").startswith("Ref Light "))


def _ensure_headers(
    worksheet,
    reading_columns: int,
    required_reading_columns: int,
    font_factory,
    fill_factory,
    alignment_factory,
) -> None:
    """Create/expand headers and move old averages when columns are added."""

    old_average_column = worksheet.max_column if worksheet.max_column >= 2 else 2
    old_reading_columns = _reading_column_count(worksheet)
    if old_reading_columns and reading_columns > old_reading_columns:
        new_average_column = reading_columns + 2
        for row_number in range(2, worksheet.max_row + 1):
            old_average = worksheet.cell(row=row_number, column=old_average_column).value
            worksheet.cell(row=row_number, column=new_average_column, value=old_average)
            worksheet.cell(row=row_number, column=old_average_column).value = None

    headers = ["Data"] + [f"Ref Light {index}" for index in range(1, reading_columns + 1)] + ["Média"]
    header_fill = fill_factory("solid", fgColor="D9EAF7")
    header_font = font_factory(bold=True)
    header_alignment = alignment_factory(horizontal="center", vertical="center")
    for column, header in enumerate(headers, start=1):
        cell = worksheet.cell(row=1, column=column, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment

    # Keep cells in newly added reading columns blank for older sessions.
    for row_number in range(2, worksheet.max_row + 1):
        for column in range(2, reading_columns + 2):
            cell = worksheet.cell(row=row_number, column=column)
            if cell.value is None:
                cell.number_format = "0.000000"
        worksheet.cell(row=row_number, column=reading_columns + 2).number_format = "0.000000"


def _column_letter(column: int) -> str:
    result = ""
    while column:
        column, remainder = divmod(column - 1, 26)
        result = chr(65 + remainder) + result
    return result
